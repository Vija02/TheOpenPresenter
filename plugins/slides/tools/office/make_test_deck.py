"""Builds an Office Online reverse-engineering fixture deck.

One feature per slide; each slide title names the feature so the GetSlide JSON
can be mapped back to a known input. Shapes are labelled A, B, C... in the order
they are animated.
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ASSETS = Path(sys.argv[1])
OUT = Path(sys.argv[2])

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
WRAP = f'<w xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:mc="{NS_MC}" xmlns:p14="{NS_P14}" xmlns:p159="{NS_P159}">{{}}</w>'

COLORS = ["E74C3C", "3498DB", "2ECC71", "F39C12", "9B59B6", "1ABC9C", "E67E22"]


def frag(xml):
    return list(etree.fromstring(WRAP.format(xml)))


# ---------------------------------------------------------------- timing XML


class Timing:
    """Emits <p:timing> for a list of click groups.

    group = {"auto": bool, "subs": [(delay_ms, [effect, ...]), ...]}
    effect = dict(kind, spid, cls, pid, sub, dur, node, grp, para, iterate, repeat, autorev, delay)
    """

    def __init__(self):
        self.n = 0
        self.bld = {}

    def nid(self):
        self.n += 1
        return self.n

    @staticmethod
    def tgt(e):
        if e.get("para") is not None:
            p = e["para"]
            return f'<p:spTgt spid="{e["spid"]}"><p:txEl><p:pRg st="{p}" end="{p}"/></p:txEl></p:spTgt>'
        return f'<p:spTgt spid="{e["spid"]}"/>'

    def cbhvr(self, e, dur, attrs=(), delay=0, additive=""):
        names = "".join(f"<p:attrName>{a}</p:attrName>" for a in attrs)
        lst = f"<p:attrNameLst>{names}</p:attrNameLst>" if attrs else ""
        add = f' additive="{additive}"' if additive else ""
        st = f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>' if delay else ""
        return (f'<p:cBhvr{add}><p:cTn id="{self.nid()}" dur="{dur}" fill="hold">{st}</p:cTn>'
                f"<p:tgtEl>{self.tgt(e)}</p:tgtEl>{lst}</p:cBhvr>")

    def set_vis(self, e, val="visible", delay=0):
        return (f'<p:set><p:cBhvr><p:cTn id="{self.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl>{self.tgt(e)}</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="{val}"/></p:to></p:set>')

    def anim(self, e, dur, attr, frm, to):
        return (f'<p:anim calcmode="lin" valueType="num">{self.cbhvr(e, dur, [attr], additive="base")}'
                f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')

    def effect_filter(self, e, dur, filt, direction="in"):
        return f'<p:animEffect transition="{direction}" filter="{filt}">{self.cbhvr(e, dur)}</p:animEffect>'

    def behaviours(self, e):
        k, d = e["kind"], e["dur"]
        if k == "appear":
            return self.set_vis(e)
        if k == "fade":
            return self.set_vis(e) + self.effect_filter(e, d, "fade")
        if k == "fly":  # from bottom
            return (self.set_vis(e) + self.anim(e, d, "ppt_x", "#ppt_x", "#ppt_x")
                    + self.anim(e, d, "ppt_y", "1+#ppt_h/2", "#ppt_y"))
        if k == "wipe":
            return self.set_vis(e) + self.effect_filter(e, d, "wipe(left)")
        if k == "zoom":
            return (self.set_vis(e) + self.anim(e, d, "ppt_w", "0", "#ppt_w")
                    + self.anim(e, d, "ppt_h", "0", "#ppt_h"))
        if k == "floatup":
            return (self.set_vis(e) + self.anim(e, d, "ppt_x", "#ppt_x", "#ppt_x")
                    + self.anim(e, d, "ppt_y", "#ppt_y+.1", "#ppt_y") + self.effect_filter(e, d, "fade"))
        if k == "fadeout":
            return self.effect_filter(e, d, "fade", "out") + self.set_vis(e, "hidden", d - 1)
        if k == "flyout":
            return (self.anim(e, d, "ppt_x", "#ppt_x", "#ppt_x") + self.anim(e, d, "ppt_y", "#ppt_y", "1+#ppt_h/2")
                    + self.set_vis(e, "hidden", d - 1))
        if k == "spin":
            return f'<p:animRot by="21600000">{self.cbhvr(e, d, ["r"])}</p:animRot>'
        if k == "grow":
            return f'<p:animScale>{self.cbhvr(e, d)}<p:by x="150000" y="150000"/></p:animScale>'
        if k == "pulse":
            return f'<p:animScale>{self.cbhvr(e, d)}<p:by x="110000" y="110000"/></p:animScale>'
        if k == "transparency":
            return (f'<p:set>{self.cbhvr(e, d, ["style.opacity"])}<p:to><p:strVal val="0.25"/></p:to></p:set>')
        if k == "fillcolor":
            return (f'<p:animClr clrSpc="rgb" dir="cw">{self.cbhvr(e, d, ["fillcolor"])}'
                    f'<p:to><a:srgbClr val="FFFF00"/></p:to></p:animClr>'
                    f'<p:set>{self.cbhvr(e, d, ["fill.type"])}<p:to><p:strVal val="solid"/></p:to></p:set>')
        if k == "path":
            return (f'<p:animMotion origin="layout" path="M 0 0 L 0.4 0.2 E" pathEditMode="relative" ptsTypes="">'
                    f'{self.cbhvr(e, d, ["ppt_x", "ppt_y"])}</p:animMotion>')
        if k == "play":
            return f'<p:cmd type="call" cmd="playFrom(0.0)">{self.cbhvr(e, d)}</p:cmd>'
        raise ValueError(k)

    def effect(self, e, node_type):
        extra = ""
        if e.get("repeat"):
            extra += f' repeatCount="{e["repeat"] * 1000}"'
        if e.get("autorev"):
            extra += ' autoRev="1"'
        grp = f' grpId="{e.get("grp", 0)}"'
        cid = self.nid()
        it = ""
        if e.get("iterate"):
            it = f'<p:iterate type="{e["iterate"]}"><p:tmPct val="10000"/></p:iterate>'
        body = self.behaviours(e)
        if e["cls"] != "mediacall" and e.get("bld", True):
            key = (e["spid"], e.get("grp", 0))
            self.bld[key] = "p" if e.get("para") is not None else self.bld.get(key, "whole")
        return (f'<p:par><p:cTn id="{cid}" presetID="{e["pid"]}" presetClass="{e["cls"]}" presetSubtype="{e.get("sub", 0)}"'
                f' fill="hold"{grp}{extra} nodeType="{node_type}"><p:stCondLst><p:cond delay="{e.get("delay", 0)}"/></p:stCondLst>'
                f"{it}<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>")

    def group(self, g, first_node):
        gid = self.nid()
        cond = '<p:cond delay="indefinite"/>'
        if g.get("auto"):
            cond += '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
        subs = ""
        for i, (delay, effects) in enumerate(g["subs"]):
            sid = self.nid()
            inner = ""
            for j, e in enumerate(effects):
                if i == 0 and j == 0:
                    nt = first_node
                elif j == 0:
                    nt = "afterEffect"
                else:
                    nt = "withEffect"
                inner += self.effect(e, e.get("node", nt))
            subs += (f'<p:par><p:cTn id="{sid}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
                     f"<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>")
        return (f'<p:par><p:cTn id="{gid}" fill="hold"><p:stCondLst>{cond}</p:stCondLst>'
                f"<p:childTnLst>{subs}</p:childTnLst></p:cTn></p:par>")

    def build(self, groups, triggers=(), videos=()):
        root, main = self.nid(), self.nid()  # ids 1, 2
        main_children = "".join(self.group(g, "afterEffect" if g.get("auto") else "clickEffect") for g in groups)
        # PowerPoint leaves the main sequence out when nothing is on click; an
        # empty one makes the file unopenable
        seqs = (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{main}" dur="indefinite" nodeType="mainSeq">'
                f"<p:childTnLst>{main_children}</p:childTnLst></p:cTn>"
                '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
                ) if main_children else ""
        for trig_spid, effects in triggers:
            tid = self.nid()
            inner = self.group({"subs": [(0, effects)]}, "clickEffect")
            # Interactive sequences start their single group at once rather than waiting for a click.
            inner = inner.replace('<p:cond delay="indefinite"/>', '<p:cond delay="0"/>', 1)
            click = f'<p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{trig_spid}"/></p:tgtEl></p:cond>'
            seqs += (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{tid}" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">'
                     f"<p:stCondLst>{click}</p:stCondLst>"
                     '<p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync>'
                     f"<p:childTnLst>{inner}</p:childTnLst></p:cTn><p:nextCondLst>{click}</p:nextCondLst></p:seq>")
        for spid in videos:
            seqs += (f'<p:video><p:cMediaNode vol="80000"><p:cTn id="{self.nid()}" fill="hold" display="0">'
                     '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>')
        bld = "".join(
            f'<p:bldP spid="{s}" grpId="{g}" animBg="1"/>' if kind == "whole" else f'<p:bldP spid="{s}" grpId="{g}" build="p"/>'
            for (s, g), kind in self.bld.items()
            if kind != "skip"
        )
        bld_lst = f"<p:bldLst>{bld}</p:bldLst>" if bld else ""
        return (f'<p:timing><p:tnLst><p:par><p:cTn id="{root}" dur="indefinite" restart="never" nodeType="tmRoot">'
                f"<p:childTnLst>{seqs}</p:childTnLst></p:cTn></p:par></p:tnLst>{bld_lst}</p:timing>")


def E(kind, spid, cls, pid, dur=500, **kw):
    return dict(kind=kind, spid=spid, cls=cls, pid=pid, dur=dur, **kw)


# ---------------------------------------------------------------- deck


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(title, notes=None):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    tb.name = "Title"
    tb.text_frame.text = title
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def box(s, label, i, shape=MSO_SHAPE.ROUNDED_RECTANGLE, x=None, y=2.2, w=2.0, h=1.6, rot=0):
    x = 0.6 + (i % 5) * 2.5 if x is None else x
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = f"Shape {label}"
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(COLORS[i % len(COLORS)])
    sp.line.fill.background()
    sp.text_frame.text = label
    sp.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    sp.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    if rot:
        sp.rotation = rot
    return sp.shape_id


def caption(s, text, y=6.4):
    tb = s.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.3), Inches(0.8))
    tb.name = "Caption"
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)


def put(s, xml):
    for el in frag(xml):
        s._element.append(el)


def transition(s, inner, spd="med", adv=""):
    put(s, f'<p:transition spd="{spd}"{adv}>{inner}</p:transition>')


def timing(s, groups, **kw):
    put(s, Timing().build(groups, **kw))


def click(*effects):
    return {"subs": [(0, list(effects))]}


# 01 baseline
s = new_slide("01 Baseline: no animation, no transition", notes="Notes for slide 01: plain text.")
box(s, "A", 0)
caption(s, "Expect: empty animLst, no transition.")

# 02-08 transitions (one per slide)
for num, label, xml, spd in [
    ("02", "Transition: Fade (slow)", "<p:fade/>", "slow"),
    ("03", "Transition: Fade through black", '<p:fade thruBlk="1"/>', "med"),
    ("04", "Transition: Push from bottom (dir=u)", '<p:push dir="u"/>', "med"),
    ("05", "Transition: Wipe down (dir=d)", '<p:wipe dir="d"/>', "fast"),
    ("06", "Transition: Split vertical out", '<p:split orient="vert" dir="out"/>', "med"),
    ("07", "Transition: Cover from right (dir=l)", '<p:cover dir="l"/>', "med"),
    ("08", "Transition: Dissolve", "<p:dissolve/>", "med"),
]:
    s = new_slide(f"{num} {label}")
    box(s, "A", 1)
    transition(s, xml, spd)

# 09 p14 explicit duration
s = new_slide("09 Transition: Fade with p14:dur=2500ms (AlternateContent)")
box(s, "A", 2)
put(s, '<mc:AlternateContent><mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="2500"><p:fade/></p:transition></mc:Choice>'
       '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')

# 10 p14-only transition (Vortex)
s = new_slide("10 Transition: Vortex (p14-only effect)")
box(s, "A", 3)
put(s, '<mc:AlternateContent><mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="2000"><p14:vortex dir="r"/></p:transition></mc:Choice>'
       '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')

# 11-12 morph pair
s = new_slide("11 Morph source: circle left, small")
box(s, "M", 4, MSO_SHAPE.OVAL, x=1.0, y=3.0, w=1.5, h=1.5)
s = new_slide("12 Transition: Morph (circle moves right + grows)")
box(s, "M", 4, MSO_SHAPE.OVAL, x=8.0, y=2.0, w=3.5, h=3.5)
put(s, '<mc:AlternateContent><mc:Choice Requires="p159"><p:transition spd="slow" p14:dur="2000"><p159:morph option="byObject"/></p:transition></mc:Choice>'
       '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')

# 13 auto-advance
s = new_slide("13 Auto-advance after 3s (advTm=3000) + fade")
box(s, "A", 5)
transition(s, "<p:fade/>", "med", ' advTm="3000"')

# 14 entrance effects, one click each
s = new_slide("14 Entrances on click: A appear, B fly-bottom, C wipe-left, D zoom, E float-up")
ids = [box(s, l, i) for i, l in enumerate("ABCDE")]
timing(s, [
    click(E("appear", ids[0], "entr", 1)),
    click(E("fly", ids[1], "entr", 2, sub=4)),
    click(E("wipe", ids[2], "entr", 22, sub=8)),
    click(E("zoom", ids[3], "entr", 23, sub=16)),
    click(E("floatup", ids[4], "entr", 42, dur=1000)),
])

# 15 with / after previous chains
s = new_slide("15 One click: A fade; B with A; C after A; D after C +1000ms delay")
ids = [box(s, l, i) for i, l in enumerate("ABCD")]
timing(s, [{"subs": [
    (0, [E("fade", ids[0], "entr", 10), E("fade", ids[1], "entr", 10, node="withEffect")]),
    (500, [E("fade", ids[2], "entr", 10, node="afterEffect")]),
    (1000, [E("fade", ids[3], "entr", 10, node="afterEffect", delay=1000)]),
]}])

# 16 auto-start after previous at slide start
s = new_slide("16 A fades in automatically (after previous, no click); B on click")
ids = [box(s, l, i) for i, l in enumerate("AB")]
timing(s, [
    {"auto": True, "subs": [(0, [E("fade", ids[0], "entr", 10, dur=1000)])]},
    click(E("fade", ids[1], "entr", 10)),
])

# 17 paragraph build
s = new_slide("17 Per-paragraph build: 4 bullets fade in, one per click")
tb = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(4))
tb.name = "Bullets"
tf = tb.text_frame
for i, t in enumerate(["First paragraph", "Second paragraph", "Third paragraph", "Fourth paragraph"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.runs[0].font.size = Pt(36)
timing(s, [click(E("fade", tb.shape_id, "entr", 10, para=i)) for i in range(4)])

# 18 letter by letter
s = new_slide("18 Fly in by letter (iterate type=lt, 10%)")
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11), Inches(1.5))
tb.name = "Letters"
tb.text_frame.text = "Letter by letter"
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(60)
timing(s, [click(E("fly", tb.shape_id, "entr", 2, sub=4, iterate="lt"))])

# 19 emphasis
s = new_slide("19 Emphasis: A spin, B grow 150%, C transparency, D pulse, E fill colour")
ids = [box(s, l, i) for i, l in enumerate("ABCDE")]
timing(s, [
    click(E("spin", ids[0], "emph", 8, dur=2000)),
    click(E("grow", ids[1], "emph", 6, dur=2000)),
    click(E("transparency", ids[2], "emph", 9, dur=2000)),
    click(E("pulse", ids[3], "emph", 26, dur=500, autorev=True)),
    click(E("fillcolor", ids[4], "emph", 1, dur=2000)),
])

# 20 exits
s = new_slide("20 Exits: A fade out, B fly out to bottom")
ids = [box(s, l, i) for i, l in enumerate("AB")]
timing(s, [
    click(E("fadeout", ids[0], "exit", 10)),
    click(E("flyout", ids[1], "exit", 2, sub=4)),
])

# 21 same shape entrance -> emphasis -> exit
s = new_slide("21 Same shape A: fade in, then spin, then fade out (3 clicks)")
a = box(s, "A", 0, x=5.5)
timing(s, [
    click(E("fade", a, "entr", 10, grp=0)),
    click(E("spin", a, "emph", 8, dur=1000, grp=1)),
    click(E("fadeout", a, "exit", 10, grp=2)),
])

# 22 motion path
s = new_slide("22 Motion path: A moves right/down along custom path")
a = box(s, "A", 0, MSO_SHAPE.OVAL, w=1.5, h=1.5)
timing(s, [click(E("path", a, "path", 0, dur=2000))])

# 23 repeat / autoreverse
s = new_slide("23 A spin repeat x3; B grow with autoReverse")
ids = [box(s, l, i) for i, l in enumerate("AB")]
timing(s, [
    click(E("spin", ids[0], "emph", 8, dur=1000, repeat=3)),
    click(E("grow", ids[1], "emph", 6, dur=1000, autorev=True)),
])

# 24 trigger
s = new_slide("24 Trigger: clicking the Button shape makes star A appear")
btn = box(s, "Button", 1, x=0.8, w=3.0, h=1.2)
star = box(s, "A", 3, MSO_SHAPE.STAR_5_POINT, x=7.0, w=2.5, h=2.5)
timing(s, [], triggers=[(btn, [E("fade", star, "entr", 10)])])

# 25 video
s = new_slide("25 Video: click 1 plays the clip; B fades on click 2")
mov = s.shapes.add_movie(str(ASSETS / "clip.mp4"), Inches(0.8), Inches(1.5), Inches(6.4), Inches(3.6),
                         poster_frame_image=str(ASSETS / "poster.png"), mime_type="video/mp4")
b = box(s, "B", 2, x=8.5)
for t in s._element.findall(f"{{{NS_P}}}timing"):  # python-pptx writes its own; replace with ours
    s._element.remove(t)
timing(s, [
    click(E("play", mov.shape_id, "mediacall", 1, dur=6000)),
    click(E("fade", b, "entr", 10)),
], videos=[mov.shape_id])

# 26 group + rotated + picture
s = new_slide("26 A: group flies in; B: rotated 30deg fades; C: picture wipes")
grp = s.shapes.add_group_shape()
grp.name = "Group A"
for i, (dx, dy) in enumerate([(0, 0), (1.2, 0.8)]):
    g = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8 + dx), Inches(2.0 + dy), Inches(1.4), Inches(1.4))
    g.fill.solid()
    g.fill.fore_color.rgb = RGBColor.from_string(COLORS[i])
    g.text_frame.text = "A"
rot = box(s, "B", 1, x=4.5, rot=30)
pic = s.shapes.add_picture(str(ASSETS / "picture.png"), Inches(8.5), Inches(1.8), Inches(3.6), Inches(2.4))
pic.name = "Picture C"
timing(s, [
    click(E("fly", grp.shape_id, "entr", 2, sub=4, bld=False)),
    click(E("fade", rot, "entr", 10)),
    click(E("wipe", pic.shape_id, "entr", 22, sub=8, bld=False)),
])

# 27 hidden slide
s = new_slide("27 Hidden slide (show=0)")
box(s, "H", 6)
s._element.set("show", "0")

# 28 durations and delays
s = new_slide("28 Timings: A fade 3000ms; B fade 250ms; C fade 500ms delay 750ms")
ids = [box(s, l, i) for i, l in enumerate("ABC")]
timing(s, [
    click(E("fade", ids[0], "entr", 10, dur=3000)),
    click(E("fade", ids[1], "entr", 10, dur=250)),
    click(E("fade", ids[2], "entr", 10, delay=750)),
], )

# 29 end
s = new_slide("29 End", notes="Notes for the last slide.")
box(s, "Z", 0)

if len(sys.argv) > 3:
    keep = {f"{int(n):02d}" for n in sys.argv[3].split(",")}
    id_lst = prs.slides._sldIdLst
    for sld_id, slide in list(zip(list(id_lst), list(prs.slides))):
        title = next(sh for sh in slide.shapes if sh.name == "Title").text_frame.text
        if title[:2] not in keep:
            prs.part.drop_rel(sld_id.rId)
            id_lst.remove(sld_id)

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
